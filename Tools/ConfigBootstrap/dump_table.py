#!/usr/bin/env python3
"""Print a Luban xlsx's header rows and data rows (read-only, for verification).

Usage: python Tools/ConfigBootstrap/dump_table.py HeroConfig.xlsx
"""

import os
import sys

import openpyxl

DATAS = r"P:\Godot-Project\GameZMXY\Configs\GameConfig\Datas"


def main():
    name = sys.argv[1]
    ws = openpyxl.load_workbook(os.path.join(DATAS, name)).active
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = [c for c in next(ws.iter_rows(min_row=r, max_row=r, values_only=True))]
        if any(v is not None for v in vals):
            print(f"{r:3d}: {vals}")


if __name__ == "__main__":
    main()
