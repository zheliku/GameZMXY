#!/usr/bin/env python3
"""一次性迁移（2026-10 架构重构 P1）：属性统一为 StatType/StatBlock，成长改为逐级表。

做了什么（全部就地改表，保留列宽和其他表；运行前已备份 Configs/GameConfig/Datas/）：
  * __enums__.xlsx：追加 Stat.StatType（16 项属性，值显式写死）、Stat.StatModifierKind（修正方式）
  * __beans__.xlsx：追加 Stat.StatBlock（16 项 int 属性）、Hero.HeroGrowthConfig、Monster.MonsterConfig
  * HeroGrowthConfig.xlsx（新）：每英雄逐级属性。数值按旧公式 Base + (Lv-1) * Grow 一次性生成，
    手感保持不变；之后由策划直接改表（可在 Excel 里写公式），代码不再出现成长公式
  * HeroConfig.xlsx：删除 Base*/Grow* 与战斗属性列，只保留身份、手感与表现
  * MonsterConfig.xlsx：Hp/Def/…/RHp 改为内嵌 Stats(StatBlock) 列组；攻击 Power 填 0
    （怪物伤害仍由招式 FlatPower 决定，与重构前一致）。记录类型改由 __beans__ 定义
    （read_schema_from_file=false，与 LevelConfig 同一写法：数据表只有多级 ##var 标题，没有 ##type）
  * ProfileConfig.xlsx（新，单行）：新建档规则 StartHeroId / StartGold
  * BattleConfig.xlsx：表尾追加 DeathRestartDelay、ClearRestartDelay（暂无结算界面时死亡/通关后自动重开的等待秒数）
  * __tables__.xlsx：注册 Hero.TbHeroGrowthConfig（list，联合索引 HeroId+Level）与 Profile.TbProfileConfig（one）

本脚本只运行一次；重复运行会检测到已迁移并退出，不会二次改写。
Run:  python Tools/ConfigBootstrap/migrate_stat_tables.py
Then: Configs/GameConfig/gen_code_bin_to_project_lazyload.bat
"""

import os
import sys

import openpyxl
from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATAS = os.path.join(ROOT, "Configs", "GameConfig", "Datas")

# 属性顺序即枚举值（0 起），也是 StatBlock 字段顺序；只允许在末尾追加。
STATS = [
    ("MaxHp", "生命上限"),
    ("MaxMp", "魔法上限"),
    ("Power", "攻击"),
    ("Def", "物防"),
    ("Mdef", "魔防"),
    ("Crit", "暴击"),
    ("Miss", "闪避"),
    ("Lucky", "幸运(加暴击倍率)"),
    ("Toughness", "韧性(减对方幸运)"),
    ("Htarget", "命中(减对方闪避)"),
    ("CritReduce", "暴击抵抗"),
    ("Ar", "破甲(减对方物防)"),
    ("Sp", "破魔(减对方魔防)"),
    ("Vampirism", "吸血(百分点,伤害转生命)"),
    ("HpRegen", "每秒回血"),
    ("MpRegen", "每秒回蓝"),
]
STAT_NAMES = [s[0] for s in STATS]

MAX_LEVEL = 55


def load(name):
    """打开已有工作簿的首张表。"""
    path = os.path.join(DATAS, name)
    wb = openpyxl.load_workbook(path)
    return path, wb, wb.active


def header_names(ws):
    """读取首行 ##var 字段名（去掉标记列与尾部空格）。"""
    names = [ws.cell(row=1, column=c).value for c in range(2, ws.max_column + 1)]
    while names and names[-1] in (None, ""):
        names.pop()
    return names


def column_of(ws, name):
    """在首行 ##var 中按字段名找列号（1 起）。"""
    for c in range(2, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == name:
            return c
    raise SystemExit(f"  [失败] 缺少列 {name}")


def data_rows(ws, first_row):
    """返回主键列非空的数据行号。"""
    return [r for r in range(first_row, ws.max_row + 1) if ws.cell(row=r, column=2).value not in (None, "")]


def rewrite_sheet(ws, rows):
    """清空整张表后按行写入（openpyxl 的 cell(value=None) 不写入，必须属性赋值）。"""
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None
    for r, cells in enumerate(rows, start=1):
        for c, val in enumerate(cells, start=1):
            ws.cell(row=r, column=c).value = val


# ---------------------------------------------------------------------------
# 1. 枚举
# ---------------------------------------------------------------------------
def add_enums():
    path, wb, ws = load("__enums__.xlsx")
    existing = {ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)}
    if "Stat.StatType" in existing:
        raise SystemExit("  [已迁移] __enums__.xlsx 已含 Stat.StatType，脚本只运行一次。")

    def append_enum(full_name, items):
        for i, (name, alias, value) in enumerate(items):
            row = [None] * 12
            if i == 0:
                row[1] = full_name
                row[2] = False
                row[3] = True
            row[7] = name
            row[8] = alias
            row[9] = value
            ws.append(row)

    append_enum("Stat.StatType", [(n, a, i) for i, (n, a) in enumerate(STATS)])
    append_enum("Stat.StatModifierKind", [
        ("Flat", "固定值", 0),
        ("PercentAdd", "同类百分比相加", 1),
        ("PercentMult", "独立百分比相乘", 2),
    ])
    wb.save(path)
    print("  __enums__.xlsx: 追加 Stat.StatType、Stat.StatModifierKind")


# ---------------------------------------------------------------------------
# 2. bean：StatBlock、英雄成长行、怪物配置记录
# ---------------------------------------------------------------------------
def append_bean(ws, full_name, comment, fields):
    """按 __beans__ 多行格式追加一个 bean：首行写全名与注释，每个字段一行。"""
    for i, (name, ftype, fcomment) in enumerate(fields):
        row = [None] * 16
        if i == 0:
            row[1] = full_name
            row[6] = comment
        row[9] = name
        row[11] = ftype
        row[13] = fcomment
        ws.append(row)


MONSTER_HEAD = ["Id", "LegacyId", "NameCn", "Desc", "EntityId", "Level"]
MONSTER_TAIL = [
    "MoveSpeed", "SightRange", "AttackDesire", "BehitCalmTime", "AddExp", "HurtSoundId", "DeathSoundId", "Gravity",
    "PatrolInterval", "PatrolIdleChance", "PatrolRadius", "AttackInterval", "AttackFirstDelay", "Rank", "SuperArmor",
    "AttackRangeSlack", "LoseTargetTime", "PaceRange",
]
# StatBlock 字段 → 旧 MonsterConfig 列；不在表里的属性填 0。
MONSTER_OLD_STATS = {
    "MaxHp": "Hp", "Def": "Def", "Mdef": "Mdef", "Crit": "Crit", "Miss": "Miss", "Lucky": "Lucky",
    "Toughness": "Toughness", "Htarget": "Htarget", "CritReduce": "CritReduce", "Ar": "Ar", "Sp": "Sp",
    "HpRegen": "RHp",
}


def add_beans(monster_types, monster_comments):
    path, wb, ws = load("__beans__.xlsx")
    existing = {ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)}
    if "Stat.StatBlock" in existing:
        raise SystemExit("  [已迁移] __beans__.xlsx 已含 Stat.StatBlock")

    append_bean(ws, "Stat.StatBlock", "一组完整属性值；英雄逐级成长行与怪物配置共用",
                [(n, "int", a) for n, a in STATS])
    append_bean(ws, "Hero.HeroGrowthConfig", "英雄逐级属性（HeroId + Level 联合主键）", [
        ("HeroId", "int", "英雄ID(对应 HeroConfig.Id)"),
        ("Level", "int", "等级(1 起连续，覆盖经验表全部等级)"),
        ("Stats", "Stat.StatBlock", "该等级的成长属性"),
    ])
    fields = [(n, monster_types[n], monster_comments[n]) for n in MONSTER_HEAD]
    fields.append(("Stats", "Stat.StatBlock", "战斗属性(攻击 Power 填 0：怪物伤害由招式 FlatPower 决定)"))
    fields += [(n, monster_types[n], monster_comments[n]) for n in MONSTER_TAIL]
    append_bean(ws, "Monster.MonsterConfig", "怪物配置（怪物无成长，Level 只用于等级压制）", fields)
    wb.save(path)
    print("  __beans__.xlsx: 追加 Stat.StatBlock、Hero.HeroGrowthConfig、Monster.MonsterConfig")


# ---------------------------------------------------------------------------
# 3. 英雄：拆出逐级成长表
# ---------------------------------------------------------------------------
HERO_KEEP = [
    "Id", "LegacyId", "NameCn", "Desc", "EntityId", "WalkSpeed", "RunSpeed", "RunDoubleTapWindow",
    "IdleEmoteDelay", "JumpSpeed", "Gravity", "JumpCountMax", "HurtSoundId", "DeathSoundId", "InputBufferTime",
]


def migrate_hero():
    path, wb, ws = load("HeroConfig.xlsx")
    names = header_names(ws)
    if "BaseHp" not in names:
        raise SystemExit("  [已迁移] HeroConfig.xlsx 已无 BaseHp 列")
    col = {name: column_of(ws, name) for name in names}

    # 先按旧公式生成逐级属性（来源：重构前 HeroEntity.ApplyExperience / GetCombatStats）。
    growth = []
    for r in data_rows(ws, 6):
        def num(key):
            value = ws.cell(row=r, column=col[key]).value
            return 0 if value in (None, "") else value

        hero_id = num("Id")
        for level in range(1, MAX_LEVEL + 1):
            g = level - 1
            stats = {
                "MaxHp": num("BaseHp") + g * num("GrowHp"),
                "MaxMp": num("BaseMp") + g * num("GrowMp"),
                "Power": num("BasePower") + g * num("GrowPower"),
                "Def": num("BaseDef") + g * num("GrowDef"),
                "Mdef": num("BaseMdef") + g * num("GrowMdef"),
                "Crit": num("Crit"), "Miss": num("Miss"), "Lucky": num("Lucky"), "Toughness": num("Toughness"),
                "Htarget": num("Htarget"), "CritReduce": num("CritReduce"), "Ar": num("Ar"), "Sp": num("Sp"),
                # 旧表这三列是 float 且全部为 0；StatBlock 统一为 int。
                "Vampirism": int(round(num("Vampirism"))),
                "HpRegen": int(round(num("RHp"))),
                "MpRegen": int(round(num("RMp"))),
            }
            growth.append((hero_id, level, stats))

    # 就地重排 HeroConfig：保留列按原顺序搬到前面，其余列清空。
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(row=r, column=1).value] + [ws.cell(row=r, column=col[n]).value for n in HERO_KEEP])
    rewrite_sheet(ws, rows)
    wb.save(path)
    print(f"  HeroConfig.xlsx: 删除属性/成长列，保留 {len(HERO_KEEP)} 列")

    # 新建逐级成长表：两级 ##var 标题（与 LevelConfig 相同），类型来自 __beans__。
    path = os.path.join(DATAS, "HeroGrowthConfig.xlsx")
    if os.path.exists(path):
        raise SystemExit("  [失败] HeroGrowthConfig.xlsx 已存在")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["##var", "HeroId", "Level", "Stats"] + [None] * (len(STATS) - 1))
    ws.append(["##var", None, None] + STAT_NAMES)
    ws.append(["##", "英雄ID(对应 HeroConfig.Id)", "等级(1 起连续)"] + [a for _, a in STATS])
    for hero_id, level, stats in growth:
        ws.append([None, hero_id, level] + [stats[n] for n in STAT_NAMES])
    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=3 + len(STATS))
    wb.save(path)
    print(f"  HeroGrowthConfig.xlsx: 新建 {len(growth)} 行（旧公式逐级展开）")


# ---------------------------------------------------------------------------
# 4. 怪物：属性列改为内嵌 StatBlock 列组
# ---------------------------------------------------------------------------
def read_monster():
    path, wb, ws = load("MonsterConfig.xlsx")
    names = header_names(ws)
    if "Hp" not in names:
        raise SystemExit("  [已迁移] MonsterConfig.xlsx 已无 Hp 列")
    col = {name: column_of(ws, name) for name in names}
    types = {name: ws.cell(row=3, column=col[name]).value for name in names}
    comments = {name: ws.cell(row=5, column=col[name]).value for name in names}
    return path, wb, ws, col, types, comments


def migrate_monster(path, wb, ws, col, comments):
    def old(r, name):
        return ws.cell(row=r, column=col[name]).value

    records = []
    for r in data_rows(ws, 6):
        stats = []
        for stat in STAT_NAMES:
            source = MONSTER_OLD_STATS.get(stat)
            stats.append(0 if source is None else int(round(old(r, source) or 0)))
        records.append([None] + [old(r, n) for n in MONSTER_HEAD] + stats + [old(r, n) for n in MONSTER_TAIL])

    stat_comments = []
    for stat, alias in STATS:
        source = MONSTER_OLD_STATS.get(stat)
        if stat == "Power":
            stat_comments.append("攻击(填 0：怪物伤害由招式 FlatPower 决定)")
        elif source is not None:
            stat_comments.append(comments[source])
        else:
            stat_comments.append(alias)

    stats_start = 2 + len(MONSTER_HEAD)
    rows = [
        ["##var"] + MONSTER_HEAD + ["Stats"] + [None] * (len(STATS) - 1) + MONSTER_TAIL,
        ["##var"] + [None] * len(MONSTER_HEAD) + STAT_NAMES + [None] * len(MONSTER_TAIL),
        ["##"] + [comments[n] for n in MONSTER_HEAD] + stat_comments + [comments[n] for n in MONSTER_TAIL],
    ] + records
    rewrite_sheet(ws, rows)
    ws.merge_cells(start_row=1, start_column=stats_start, end_row=1, end_column=stats_start + len(STATS) - 1)
    wb.save(path)
    print(f"  MonsterConfig.xlsx: 属性改为 Stats 列组，{len(records)} 行")


# ---------------------------------------------------------------------------
# 5. 档案配置、战斗常数与表注册
# ---------------------------------------------------------------------------
def add_profile_config():
    path = os.path.join(DATAS, "ProfileConfig.xlsx")
    if os.path.exists(path):
        raise SystemExit("  [失败] ProfileConfig.xlsx 已存在")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["##var", "NameCn", "Desc", "StartHeroId", "StartGold"])
    ws.append(["##var"])
    ws.append(["##type", "string", "string", "int", "int"])
    ws.append(["##group"])
    ws.append(["##", "中文名", "描述", "新档的出战英雄(HeroConfig.Id)", "新档初始金币"])
    ws.append([None, "建档规则", "新建存档时的初始档案", 1, 0])
    wb.save(path)
    print("  ProfileConfig.xlsx: 新建（单行）")


def add_battle_constant():
    path, wb, ws = load("BattleConfig.xlsx")
    names = header_names(ws)
    if "DeathRestartDelay" in names:
        raise SystemExit("  [已迁移] BattleConfig.xlsx 已含 DeathRestartDelay")
    c = len(names) + 2
    ws.cell(row=1, column=c).value = "DeathRestartDelay"
    ws.cell(row=3, column=c).value = "float"
    ws.cell(row=5, column=c).value = "英雄死亡后自动重开本关的等待秒数(含死亡动画 1.1 秒)"
    ws.cell(row=6, column=c).value = 2.5
    ws.cell(row=1, column=c + 1).value = "ClearRestartDelay"
    ws.cell(row=3, column=c + 1).value = "float"
    ws.cell(row=5, column=c + 1).value = "通关后自动重开本关的等待秒数(暂无结算界面时使用)"
    ws.cell(row=6, column=c + 1).value = 3.0
    wb.save(path)
    print("  BattleConfig.xlsx: 表尾追加 DeathRestartDelay、ClearRestartDelay")


def update_table_registry():
    path, wb, ws = load("__tables__.xlsx")
    existing = {ws.cell(row=r, column=2).value: r for r in range(4, ws.max_row + 1)}
    for name in ("Hero.TbHeroGrowthConfig", "Profile.TbProfileConfig"):
        if name in existing:
            raise SystemExit(f"  [已迁移] __tables__.xlsx 已含 {name}")
    # MonsterConfig 改由 __beans__ 定义记录类型（数据表只剩多级 ##var 标题）。
    monster_row = existing["Monster.TbMonsterConfig"]
    ws.cell(row=monster_row, column=4).value = False
    ws.cell(row=monster_row, column=6).value = "Id"
    ws.cell(row=monster_row, column=7).value = "map"
    ws.append([None, "Hero.TbHeroGrowthConfig", "HeroGrowthConfig", False, "HeroGrowthConfig.xlsx", "HeroId+Level",
               "list", None, None, None, None])
    ws.append([None, "Profile.TbProfileConfig", "ProfileConfig", True, "ProfileConfig.xlsx", None, "one",
               None, None, None, None])
    wb.save(path)
    print("  __tables__.xlsx: 注册 TbHeroGrowthConfig、TbProfileConfig；TbMonsterConfig 改为 bean 定义")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    monster = read_monster()
    add_enums()
    add_beans(monster[4], monster[5])
    migrate_hero()
    migrate_monster(monster[0], monster[1], monster[2], monster[3], monster[5])
    add_profile_config()
    add_battle_constant()
    update_table_registry()
    print("完成。下一步：Configs/GameConfig/gen_code_bin_to_project_lazyload.bat")


if __name__ == "__main__":
    main()
