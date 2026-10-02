#!/usr/bin/env python3
"""M6 关卡流程表变更（一次性，可重跑、可追溯）。

依据 Docs 计划（.kilo/plans/1790923848521-m6-level-flow-ui-plan.md §2）：

  1) LevelSpawnConfig: 第 2~4 波 MonsterId 2/3 → 1（Monster_2/3 未迁，垂直切片全用花果山猴子；
     定点改单元格，不动其他数据）
  2) LevelConfig: BgmPath(string) → BgmSoundId(Sound.SoundId)（LegacyAssetMap 定稿设计），
     花果山行填 Level1Bgm —— 结构变更，force_rows 重写唯一数据行
  3) BattleConfig: 表尾追加 WsMax 列（旧项目无双满值 100），回填不覆盖手工数值
  4) SoundConfig: 追加 Level1Bgm 行（Id=7，Music 组，指向 Audios/BGM/level_1.mp3）
  5) __enums__.xlsx: Sound.SoundId 追加 Level1Bgm=7（值显式写死，见 Configs/AGENTS.md）

Run:  uv run --with openpyxl python Tools/ConfigBootstrap/build_m6_tables.py
Then: Configs/GameConfig/gen_code_bin_to_project_lazyload.bat  →  dotnet build

同时 build_m2_tables.py 的种子已同步到 M6 后形态（防止重跑 M2 脚本时 schema/枚举回归）。
"""

import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from build_m2_tables import DATAS, new_table  # noqa: E402


def patch_level_spawn():
    """MonsterId 2/3 → 1：定点改单元格。改完顺带更新列注释（2/3 已不再引用）。"""
    path = os.path.join(DATAS, "LevelSpawnConfig.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    changed = 0
    for r in range(6, ws.max_row + 1):
        cell = ws.cell(row=r, column=6)  # MonsterId 列（1 起：A 空 + Id..Y）
        if cell.value in (2, 3):
            cell.value = 1
            changed += 1
    ws.cell(row=5, column=6).value = (
        "怪物ID(MonsterConfig.Id;M6 垂直切片全用花果山猴子,旧 Monster_2/3 迁入后改回)"
    )
    wb.save(path)
    print(f"  LevelSpawnConfig.xlsx: {changed} 行 MonsterId 2/3 → 1")


def patch_level_config():
    """BgmPath → BgmSoundId：结构变更（中间列改名换类型），force_rows 重写唯一数据行。"""
    fields = [
        ("Id", "int", "关卡ID"),
        ("LegacyId", "int", "旧关卡编号"),
        ("NameCn", "string", "中文名"),
        ("Desc", "string", "描述"),
        ("ScenePath", "string", "关卡场景路径"),
        ("BgmSoundId", "Sound.SoundId", "关卡背景音乐(SoundConfig;None=无;触发者在此,按 Configs/AGENTS.md 音效定稿)"),
        ("MaxAlive", "int", "场上怪物上限(旧项目全局硬编码6)"),
        ("SpawnInterval", "float", "补怪间隔秒(旧设置默认1.2)"),
    ]
    rows = [
        (1, 1, "花果山", "第一关", "res://TheGame/Scenes/Level_1.tscn", "Level1Bgm", 6, 1.2),
    ]
    new_table("LevelConfig.xlsx", fields, rows, force_rows=True)


def patch_battle_config():
    """表尾追加 WsMax（安全回填：只补空单元格，已有值一律不动）。"""
    fields = [
        ("NameCn", "string", "中文名"),
        ("Desc", "string", "描述"),
        ("DefK_Hero", "float", "英雄作防守方时物防减伤K: x/(x+K)"),
        ("MdefK_Hero", "float", "英雄作防守方时魔防减伤K"),
        ("MissK_Hero", "float", "英雄作防守方时被命中K: Miss/(Miss+K)"),
        ("DefK_Monster", "float", "怪物作防守方时物防减伤K"),
        ("MdefK_Monster", "float", "怪物作防守方时魔防减伤K"),
        ("MissK_Monster", "float", "怪物作防守方时被命中K"),
        ("CritK", "float", "暴击率基数: Crit/(Crit+K)"),
        ("LuckyK_HeroAtk", "float", "英雄攻击时暴击倍率幸运K"),
        ("LuckyK_MonsterAtk", "float", "怪物攻击时暴击倍率幸运K"),
        ("CritBase", "float", "暴击倍率基数: CritBase + Lucky/(Lucky+K)"),
        ("LvMissCoef_HeroDef", "float", "等级压制-英雄防守闪避系数"),
        ("LvCritCoef_HeroDef", "float", "等级压制-英雄防守暴击系数"),
        ("LvMissCoef_MonsterDef", "float", "等级压制-怪物防守闪避系数"),
        ("LvCritCoef_MonsterDef", "float", "等级压制-怪物防守暴击系数"),
        ("LvLuckyCoef", "float", "等级压制-幸运系数"),
        ("LvDamageCoef", "float", "等级压制-每级伤害系数"),
        ("LvMissCap_HeroDef", "float", "等级压制-英雄防守闪避系数封顶"),
        ("LvMissCap_MonsterDef", "float", "等级压制-怪物防守闪避系数封顶"),
        ("LvCritCap", "float", "等级压制-暴击系数封顶(人怪两侧一致)"),
        ("LvLuckyCap", "float", "等级压制-幸运系数封顶(人怪两侧一致)"),
        ("LvDamageCapLv_HeroDef", "int", "等级压制-英雄防守时伤害最多按几级算"),
        ("LvDamageCapLv_MonsterDef", "int", "等级压制-怪物防守时伤害最多按几级算"),
        ("KnockbackScaleX_HeroDef", "float", "英雄被击退横向换算: px/s = Knockback.X × 本值"),
        ("KnockbackScaleX_MonsterDef", "float", "怪物被击退横向换算: px/s = Knockback.X × 本值"),
        ("KnockbackScaleY", "float", "击退纵向换算: px/s = Knockback.Y × 本值(人怪一致)"),
        # M6 表尾追加：无双值上限（旧项目满值 100 放无双技能；M6 无技能，仅 HUD 显示）
        ("WsMax", "int", "无双值上限(旧项目满值 100;M6 供 HUD 无双条满值)"),
    ]
    # 与现有行完全同值 + 表尾 WsMax=100（只回填空单元格）
    rows = [
        ("战斗常数", "沿用旧项目人怪两侧不一致的基数,集中在此处",
         250, 250, 100, 100, 100, 70, 100, 100, 50, 2,
         0.03, 0.07, 0.04, 0.11, 0.07, 0.05,
         1, 0.9, 1, 0.7, 5, 2, 25, 30, 15, 100),
    ]
    new_table("BattleConfig.xlsx", fields, rows)


def patch_sound_config():
    """追加 Level1Bgm 行（表已有数据 → 直接 append，不触碰现有行）。"""
    path = os.path.join(DATAS, "SoundConfig.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == 7:
            print("  SoundConfig.xlsx: 已有 Id=7 行，跳过")
            return
    ws.append([None, 7, "Level1Bgm", "花果山BGM", "花果山关卡背景音乐(关卡进场播放)",
               "1_music.mp3", "Music", "res://TheGame/Audios/BGM/level_1.mp3"])
    wb.save(path)
    print("  SoundConfig.xlsx: 追加 Level1Bgm 行（Id=7）")


def patch_enums():
    """Sound.SoundId 追加 Level1Bgm=7（在最后一个 SoundId 条目后追加一行）。"""
    path = os.path.join(DATAS, "__enums__.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    # 找 SoundId 块的最后一条（值列为第 10 列，枚举名列第 8 列）
    last = None
    for r in range(4, ws.max_row + 1):
        if ws.cell(row=r, column=8).value == "WukongImpact":
            last = r
    if last is None:
        raise SystemExit("  [拒绝] __enums__.xlsx 里没找到 Sound.SoundId 的 WukongImpact 条目")
    if last < ws.max_row and ws.cell(row=last + 1, column=8).value == "Level1Bgm":
        print("  __enums__.xlsx: SoundId 已有 Level1Bgm，跳过")
        return
    row = [None] * 12
    row[7] = "Level1Bgm"
    row[8] = "花果山BGM"
    row[9] = 7
    row[10] = "花果山关卡背景音乐"
    ws.append(row)
    wb.save(path)
    print("  __enums__.xlsx: Sound.SoundId + Level1Bgm = 7")


if __name__ == "__main__":
    print("M6 table patches:")
    patch_level_spawn()
    patch_level_config()
    patch_battle_config()
    patch_sound_config()
    patch_enums()
    print("done.")
